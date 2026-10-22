"""Arma el catálogo de seguros a partir del Excel estandarizado del cliente ("Apoyo_Seguros_estandarizado.xlsx").

Uso:
    python scripts/prompt_bench/extract_catalog.py "<ruta>/Apoyo_Seguros_estandarizado.xlsx"

Escribe data/catalog/products.json (data/ está en .gitignore: es material del cliente, no va al repo).

El Excel tiene una hoja por tema y la columna "ID" (P01, P02, ...) relaciona todas las hojas:
- Productos: una fila por producto retenible (compañía, pólizas, prima, capital, edades, devolución, retracto, ...).
- Coberturas, Asistencias, Primas, Siniestros, Indemnizacion_IT, Exclusiones, Tasas_Credito, Edad_Max_Plazo: filas
  por producto.
- No_Retenibles y No_Renunciables: seguros que no tienen ficha, solo categoría y gestión.
- Medios_Pago, Procedimientos, Equivalencias, Observaciones, Comparativa_Mercado: datos de referencia.
Exclusiones, tasas y edad máxima por plazo se guardan en la ficha (catálogo v1.1.0) pero no van al prompt de
validación (render.py no los lee).

El match con la llamada es por "Nombre en sistema" (el nombre que llega en los metadatos de BancoEstado). Una celda
puede traer varios nombres separados por " / ". Se comparan sin tildes ni mayúsculas.
"""

import json
import sys
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "data" / "catalog" / "products.json"

CHANNELS = {"MUS", "SALESFORCE"}
PREMIUM_TYPES = {"mensual": "mensual", "anual": "anual", "crédito": "credito", "credito": "credito"}
NON_CANCELLABLE = {
    "derivar a sucursal": "NO_RENUNCIABLE_SUCURSAL",
    "lo paga el banco": "NO_RENUNCIABLE_COSTO_BANCO",
    "subsidiado minvu (obligatorio)": "NO_RENUNCIABLE_SUBSIDIADO",
    "obligatorio": "NO_RENUNCIABLE_OBLIGATORIO",
}


def normalize(name: str) -> str:
    text = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode()
    return " ".join(text.upper().split())


def clean(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def rows(wb, sheet: str) -> list[dict]:
    ws = wb[sheet]
    header = [clean(c.value) for c in ws[1]]
    out = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in values):
            continue
        out.append({h: clean(v) for h, v in zip(header, values) if h})
    return out


def by_id(items: list[dict]) -> dict[str, list[dict]]:
    grouped: dict[str, list[dict]] = {}
    for item in items:
        grouped.setdefault(item.get("ID"), []).append(item)
    return grouped


def main() -> None:
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    coverages = by_id(rows(wb, "Coberturas"))
    assistances = by_id(rows(wb, "Asistencias"))
    premiums = by_id(rows(wb, "Primas"))
    claims = by_id(rows(wb, "Siniestros"))
    it_table = by_id(rows(wb, "Indemnizacion_IT"))
    exclusions = by_id(rows(wb, "Exclusiones"))
    credit_rates = by_id(rows(wb, "Tasas_Credito"))
    max_ages = by_id(rows(wb, "Edad_Max_Plazo"))
    real_names = {normalize(r["Nombre en sistema"]): r["Nombre real"] for r in rows(wb, "Equivalencias")}
    observations: dict[str, list[str]] = {}
    for obs in rows(wb, "Observaciones"):
        for pid in (obs.get("ID") or "").split("/"):
            observations.setdefault(pid.strip(), []).append(f"{obs['Tema']}: {obs['Observación']}")

    products = []
    for p in rows(wb, "Productos"):
        pid = p["ID"]
        # En las fichas HDI la columna trae el canal ("MUS", "Salesforce") en vez del nombre en sistema, o viene vacía:
        # esos seguros se encuentran por nombre de producto (catalog_match.py).
        names = [n.strip() for n in (p["Nombre en sistema"] or "").split(" / ") if n.strip()]
        channels = [n for n in names if normalize(n) in CHANNELS]
        system_names = [n for n in names if normalize(n) not in CHANNELS]
        products.append({
            "id": pid,
            "system_names": system_names,
            "channels": channels,
            "source_sheet": p.get("Hoja original"),
            "name": p["Producto"],
            "insurer": p["Compañía"],
            "category": "RETENIBLE" if (p["Grupo"] or "").lower() == "retenible" else p["Grupo"],
            "premium_type": PREMIUM_TYPES.get((p["Tipo"] or "").lower(), "sin_dato"),
            "premium_mode": p["Prima / modalidad"],
            "policies": p["Pólizas"],
            "insured_capital": p["Capital asegurado"],
            "min_entry_age": p["Edad mín. ingreso"],
            "max_entry_age": p["Edad máx. ingreso"],
            "max_stay_age": p["Edad máx. permanencia"],
            "claim_deadline_days": p["Plazo denuncio siniestro (días)"],
            "premium_refund": p["Devolución de primas"],
            "withdrawal": p["Retracto"],
            "coverage_start": p["Inicio de cobertura"],
            "solidarity_support": p["Apoyo solidario"],
            "assistance_phone": p["Fono asistencias"],
            "notes": p["Notas"],
            "coverages": [
                {k: c.get(v) for k, v in (("name", "Cobertura"), ("amount", "Monto / capital"), ("detail", "Detalle"), ("beneficiary", "Beneficiario"))}
                for c in coverages.get(pid, [])
            ],
            "assistances": [
                {k: a.get(v) for k, v in (("name", "Asistencia"), ("detail", "Detalle"), ("events", "Eventos"), ("cap", "Tope"), ("phone", "Fono"))}
                for a in assistances.get(pid, [])
            ],
            "premiums": [
                {k: r.get(v) for k, v in (("plan", "Plan"), ("insured", "Asegurados"), ("age_band", "Tramo de edad"), ("amount_insured", "Monto asegurado"),
                                          ("premium_uf", "Prima UF"), ("premium_clp", "Prima $ (aprox.)"), ("periodicity", "Periodicidad"),
                                          ("policy_table", "Póliza / tabla"), ("note", "Nota"))}
                for r in premiums.get(pid, [])
            ],
            "claims": [
                {k: c.get(v) for k, v in (("coverage", "Cobertura"), ("deadline", "Plazo para denunciar"), ("documents", "Documentos requeridos"), ("channel", "Canal de denuncio"),
                                          ("coverage_term", "Plazo de cobertura"))}
                for c in claims.get(pid, [])
            ],
            "temporary_disability_table": [
                {"from_day": r.get("Días de incapacidad desde"), "to_day": r.get("Días hasta"), "amount": r.get("Monto a indemnizar")} for r in it_table.get(pid, [])
            ],
            "observations": observations.get(pid, []),
            "exclusions": [{"type": r.get("Tipo"), "detail": r.get("Detalle")} for r in exclusions.get(pid, [])],
            "credit_rates": [
                {k: r.get(v) for k, v in (("coverage", "Cobertura"), ("term_from", "Plazo desde (meses)"), ("term_to", "Plazo hasta (meses)"),
                                          ("net_rate", "Tasa neta ‰"), ("vat", "IVA ‰"), ("gross_rate", "Tasa bruta ‰"), ("note", "Nota"))}
                for r in credit_rates.get(pid, [])
            ],
            "max_age_by_term": [
                {"term_from": r.get("Plazo desde (meses)"), "term_to": r.get("Plazo hasta (meses)"), "max_age": r.get("Edad máx. al otorgamiento")}
                for r in max_ages.get(pid, [])
            ],
        })

    # Seguros sin ficha: no renunciables (mandan sobre no retenibles, ej. el SOAP está en ambas listas) y no retenibles.
    listed = {normalize(n) for p in products for n in p["system_names"]}
    for r in rows(wb, "No_Renunciables"):
        name = r["Nombre en sistema"]
        listed.add(normalize(name))
        products.append({
            "id": None, "system_names": [name], "name": real_names.get(normalize(name), name), "insurer": None,
            "category": NON_CANCELLABLE.get((r["Categoría"] or "").lower(), "NO_RENUNCIABLE"),
            "category_detail": r["Gestión"], "premium_type": "sin_dato",
        })
    for r in rows(wb, "No_Retenibles"):
        name = r["Nombre en sistema"]
        if normalize(name) in listed:
            continue
        products.append({
            "id": None, "system_names": [name], "name": name, "insurer": r["Compañía"],
            "category": "NO_RETENIBLE", "premium_type": "sin_dato",
        })

    catalog = {
        "source": Path(sys.argv[1]).name,
        "products": products,
        "payment_methods": {r["Código"]: r["Medio de pago"] + (f" ({r['Detalle']})" if r.get("Detalle") else "") for r in rows(wb, "Medios_Pago")},
        "procedures": rows(wb, "Procedimientos"),
        # Comparación con otros bancos por segmento: no es de un seguro puntual (dato de referencia del catálogo).
        "market_comparison": rows(wb, "Comparativa_Mercado"),
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(catalog, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    with_sheet = sum(1 for p in products if p["id"])
    print(f"{len(products)} seguros ({with_sheet} con ficha) -> {OUT}")


if __name__ == "__main__":
    main()
